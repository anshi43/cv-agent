"""Markdown -> .docx (ATS-safe: single column, no tables, standard headings)."""
from __future__ import annotations

import io
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _add_runs(par, text: str):
    """Very small inline markdown renderer: **bold**, *italic*, `code`."""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part)


def markdown_to_docx(md: str, base_font: str = "Calibri", size: int = 10) -> bytes:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = base_font
    style.font.size = Pt(size)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = docx.shared.Cm(1.6)
        s.left_margin = s.right_margin = docx.shared.Cm(1.8)

    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            _add_runs(p, line[2:])
            p.runs and setattr(p.runs[0].font, "italic", True)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                r.bold = True
                r.font.size = Pt(size + 8)
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(text.upper())
                r.bold = True
                r.font.size = Pt(size + 1)
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(0)
                _add_runs(p, text)
                for r in p.runs:
                    r.bold = True
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _add_runs(p, re.sub(r"^\s*[-*]\s+", "", line))
            continue
        if re.match(r"^-{3,}$", line.strip()):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        _add_runs(p, line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def markdown_to_txt(md: str) -> str:
    t = re.sub(r"^#{1,6}\s*", "", md, flags=re.MULTILINE)
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return t
