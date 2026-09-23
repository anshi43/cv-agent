"""Document loading: PDF / DOCX / TXT -> plain text."""
from __future__ import annotations

import io
import re


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # strip weird bullet glyphs -> "- "
    text = re.sub(r"^[\s]*[•▪◦‣·–—\*]\s*", "- ", text, flags=re.MULTILINE)
    return text.strip()


def from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return _clean("\n".join(pages))


def from_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return _clean("\n".join(parts))


def from_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return _clean(data.decode(enc))
        except UnicodeDecodeError:
            continue
    return _clean(data.decode("utf-8", errors="ignore"))


def load(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return from_pdf(data)
    if name.endswith(".docx"):
        return from_docx(data)
    if name.endswith((".txt", ".md", ".markdown")):
        return from_txt(data)
    if name.endswith(".doc"):
        raise ValueError("Legacy .doc is not supported — save as .docx or PDF first.")
    # best effort
    return from_txt(data)
